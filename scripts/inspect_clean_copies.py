import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

for filename in ['Clean Copies of UNILAG Communication Review Journal Articles (1).docx', 'Clean Copies of UNILAG Communication Review Journal Articles .docx']:
    try:
        doc = docx.Document(filename)
        print(f"=== {filename} ===")
        print("Total paragraphs:", len(doc.paragraphs))
        print("Total tables:", len(doc.tables))
        print("Total shapes:", len(doc.inline_shapes))
        
        # Check first 30 paragraphs
        for i in range(min(30, len(doc.paragraphs))):
            t = doc.paragraphs[i].text.strip()
            if t:
                print(f"P{i}: {t[:120]}")
        break
    except Exception as e:
        print(f"Error reading {filename}: {e}")
