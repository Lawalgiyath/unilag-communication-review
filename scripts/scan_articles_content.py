import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('Clean Copies of UNILAG Communication Review Journal Articles .docx')

# Let's find all article titles in the document
articles = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    # Check if text is all uppercase and looks like a title
    if txt.isupper() and len(txt) > 20 and not txt.startswith('TABLE') and not txt.startswith('FIGURE') and not txt.startswith('REFERENCES') and not txt.startswith('ABSTRACT') and not txt.startswith('INTRODUCTION'):
        articles.append((i, txt))

print(f"Found {len(articles)} articles:")
for idx, (p_idx, title) in enumerate(articles):
    print(f"\nArticle {idx+1} at P{p_idx}: {title}")

# Check tables in document
print(f"\n=== Total Tables: {len(doc.tables)} ===")
for i, table in enumerate(doc.tables):
    # find where this table is or sample its text
    headers = [cell.text.strip().replace('\n', ' ') for cell in table.rows[0].cells[:4]] if len(table.rows) > 0 else []
    print(f"Table {i+1} ({len(table.rows)} rows x {len(table.columns)} cols): {' | '.join(headers)[:100]}")

# Check figures / captions in paragraphs
print("\n=== Searching for Figure Captions ===")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.lower().startswith('figure') or txt.lower().startswith('fig.') or 'figure ' in txt.lower()[:15]:
        print(f"P{i}: {txt[:100]}")
